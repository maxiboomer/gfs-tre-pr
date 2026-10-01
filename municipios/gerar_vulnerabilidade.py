#!/usr/bin/env python3
"""Gera municipios/vulnerabilidade.json (7 campos) a partir do PANORAMA ESTADUAL.xlsx
e valida contra os nomes do artefato (D.mun).
Colunas (LOCAIS_DE_VOTACAO_BD): 3=MUNICIPIO, 12=ELEITORES, 14=STATUS_LOCAL,
20=ALAGAMENTO_NIVEL, 26=VENDAVAIS_NIVEL, 32=PONTES_NIVEL, 38=ENERGIA_NIVEL."""
import openpyxl, json, re, unicodedata, hashlib

ORD = {'Baixo':0,'Não tratar':0,'Médio':1,'Alto':2,'Muito Alto':3,'Crítico':4}
C_NOME, C_ELE, C_STAT = 3, 12, 14
C_RISCO = {'en':38,'al':20,'ven':26,'pon':32}

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFD',s) if unicodedata.category(c)!='Mn').lower()

h = open('/root/artefato_final.html',encoding='utf-8').read()
D = json.loads(re.search(r'const D=(\{.*?\});\n', h, re.DOTALL).group(1))
art = {m['n'] for m in D['mun']}
by_norm = {}
for n in art: by_norm.setdefault(norm(n), []).append(n)
FIX = {'MUNHOZ DE MELLO':'Munhoz de Melo','MUNHOZ DE MELO':'Munhoz de Melo'}

wb = openpyxl.load_workbook('/root/municipios/PANORAMA ESTADUAL.xlsx', data_only=True)
ws = wb['LOCAIS_DE_VOTACAO_BD']
mun = {}
for r in range(2, ws.max_row+1):
    nome = (ws.cell(r, C_NOME).value or '').strip().upper()
    if not nome: continue
    el = int(ws.cell(r, C_ELE).value or 0)
    d = mun.setdefault(nome, {'el':0,'loc':0,'en':0,'al':0,'ven':0,'pon':0,'loc_in':0})
    d['el'] += el; d['loc'] += 1
    st = str(ws.cell(r, C_STAT).value or '')
    if 'INAB' in st.upper():
        d['loc_in'] += 1
    for slot,cn in C_RISCO.items():
        v = ws.cell(r, cn).value
        if v is not None:
            s = str(v).strip()
            if ORD.get(s) is not None and ORD[s] > d[slot]:
                d[slot] = ORD[s]

out = {}
erros = []
for nome,d in mun.items():
    alvo = FIX.get(nome)
    if not alvo:
        key = norm(nome)
        alvo = by_norm[key][0] if key in by_norm else None
    if not alvo:
        erros.append(nome); continue
    out[alvo] = [d['en'],d['al'],d['el'],d['loc'],d['loc_in'],d['ven'],d['pon']]
for n in art:
    if n not in out: out[n] = [0,0,0,0,0,0,0]

p = '/root/gfs-tre-pr/municipios/vulnerabilidade.json'
open(p,'w').write(json.dumps(out, ensure_ascii=False, separators=(',',':')))
print("municipios:", len(out), "| sem-match:", erros)
print("EN>=2:", sum(1 for v in out.values() if v[0]>=2),
      "| AL>=2:", sum(1 for v in out.values() if v[1]>=2),
      "| VEN>=2:", sum(1 for v in out.values() if v[5]>=2),
      "| PONT>=2:", sum(1 for v in out.values() if v[6]>=2),
      "| loc_in>0:", sum(1 for v in out.values() if v[4]>0))
print("sha:", hashlib.sha256(open(p,'rb').read()).hexdigest()[:12])
